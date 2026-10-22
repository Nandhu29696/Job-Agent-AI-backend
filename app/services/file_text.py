"""Extract plain text from uploaded resume files (PDF / DOCX / TXT)."""
import io

from fastapi import HTTPException, UploadFile

MAX_UPLOAD_BYTES = 5 * 1024 * 1024


async def extract_text(file: UploadFile) -> str:
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 5 MB)")
    name = (file.filename or "").lower()

    try:
        if name.endswith(".pdf") or raw[:4] == b"%PDF":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            if len(reader.pages) > 15:
                raise HTTPException(status_code=400, detail="PDF has too many pages (max 15)")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx") or raw[:2] == b"PK":
            import docx

            document = docx.Document(io.BytesIO(raw))
            lines = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(lines)
        elif name.endswith((".txt", ".md")):
            text = raw.decode("utf-8", errors="ignore")
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Upload PDF, DOCX or TXT.")
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001 — malformed files
        raise HTTPException(status_code=400, detail="Could not read this file. Try another format.") from exc

    text = text.strip()
    if len(text) < 50:
        raise HTTPException(status_code=400,
                            detail="Could not extract text — scanned/image PDFs are not supported.")
    return text[:20000]
